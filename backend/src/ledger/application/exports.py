"""CSV／XLSX 导出（需求 §5，ACC-17）：与 Dashboard 同一筛选范围；防公式注入；短期下载链接。"""

from __future__ import annotations

import csv
import io
from datetime import timedelta
from decimal import Decimal
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook

from ledger.adapters.blobs import BlobStore
from ledger.adapters.dynamo import keys
from ledger.domain import dashboard as dash
from ledger.domain.authz import Actor, require_member
from ledger.domain.errors import NotFound
from ledger.domain.money import format_minor

from . import actions
from .actions import Built, Scope
from .context import AppContext
from .entries import _consistent_month, _Env, _parse_filters
from .guard import guard_member, load_membership

DOWNLOAD_TTL_S = 300
TYPE_ZH = {
    "expense": "支出",
    "income": "收入",
    "refund": "退款",
    "internal_transfer": "内部转账",
    "exchange": "换汇",
    "card_repayment": "信用卡还款",
    "receivable": "往来款",
}
METHOD_ZH = {"credit_card": "信用卡", "debit_card": "借记卡", "cash": "现金"}
HEADER = [
    "日期",
    "类型",
    "一级分类",
    "二级分类",
    "币种",
    "原始金额",
    "合NZD",
    "合CNY",
    "汇率日期",
    "备注",
    "消费/收款方式",
    "录入人",
    "是否计入收支",
]
_RISKY = ("=", "+", "-", "@", "\t", "\r", "\n")


def safe_text(v: str) -> str:
    """防止电子表格把文本当公式执行（OWASP CSV Injection）。"""
    return "'" + v if v.startswith(_RISKY) else v


def _rows(env: _Env, entries: list[Any]) -> list[list[Any]]:
    out = []
    for e in entries:
        c = env.catalog.resolve(e.leaf_category_id)
        digits = env.currencies[e.currency].minor_digits if e.currency in env.currencies else 2
        name, left = env.names.get(e.created_by, ("", True))
        out.append(
            [
                e.business_date.isoformat(),
                TYPE_ZH[e.type],
                c.parent_name,
                c.leaf_name,
                e.currency,
                Decimal(format_minor(e.amount_minor, digits)),
                Decimal(format_minor(e.nzd_minor, 2)),
                Decimal(format_minor(e.cny_minor, 2)),
                e.fx_snapshot.effective_date.isoformat(),
                e.note,
                METHOD_ZH.get(e.payment_method or "", ""),
                name + ("（已离开）" if left else ""),
                "是" if e.counts_in_stats else "否",
            ]
        )
    return out


TEXT_COLS = {1, 2, 3, 4, 9, 10, 11, 12}


def to_csv(rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(HEADER)
    for r in rows:
        w.writerow([safe_text(str(v)) if i in TEXT_COLS else str(v) for i, v in enumerate(r)])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rows: list[list[Any]], summary: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "账目"
    ws.append(HEADER)
    for r in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
        for i in TEXT_COLS:
            cell = ws.cell(row=ws.max_row, column=i + 1)
            cell.data_type = "s"  # 强制文本：以 = 开头的备注不会成为公式
        for col in (6, 7, 8):
            ws.cell(row=ws.max_row, column=col).number_format = "#,##0.00"
    s = wb.create_sheet("汇总")
    t = summary["totals"]
    s.append(["项目", "合NZD", "合CNY"])
    for label, k in (
        ("总收入", "income"),
        ("消费总额", "expense_gross"),
        ("退款总额", "refund"),
        ("净支出", "net_expense"),
        ("结余", "balance"),
    ):
        s.append([label, float(t[k]["nzd"]), float(t[k]["cny"])])
    s.append(["笔数", summary["entry_count"], ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_export(
    ctx: AppContext,
    blobs: BlobStore,
    actor: Actor,
    fid: str,
    key: str,
    *,
    fmt: str,
    month: str,
    filters: dict[str, Any] | None,
) -> dict[str, Any]:
    body = {"fmt": fmt, "month": month, "filters": filters or {}}

    def build() -> Built[dict[str, Any]]:
        env = _Env(ctx, actor, fid)
        f = _parse_filters(filters)
        entries, _ = _consistent_month(ctx, fid, month)
        matched = [e for e in entries if e.state == "active" and dash.matches(e, f, env.catalog)]
        rows = _rows(env, matched)
        summary = dash.compute(matched, env.catalog)
        data = to_csv(rows) if fmt == "csv" else to_xlsx(rows, summary)
        jid = ctx.ids()
        ext = "csv" if fmt == "csv" else "xlsx"
        blob = f"exports/{fid}/{jid}.{ext}"
        ctype = (
            "text/csv; charset=utf-8"
            if fmt == "csv"
            else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        version = blobs.put(blob, data, ctype)
        now = ctx.clock()
        item = {
            "PK": keys.family(fid),
            "SK": f"EXPORT#{jid}",
            "type": "export",
            "job_id": jid,
            "status": "ready",
            "format": fmt,
            "month": month,
            "filters": f.to_dict(),
            "row_count": len(rows),
            "s3_key": blob,
            "s3_version_id": version,
            "created_by": actor.user_id,
            "created_at": keys.ts(now),
            "ttl": int((now + timedelta(days=7)).timestamp()),
        }
        tx = ctx.store.tx()
        guard_member(tx, actor, env.m)
        tx.put_new(item)
        return Built(tx, {"job_id": jid}, [{"object_type": "export", "object_id": jid}])

    out = actions.run(
        ctx,
        actor,
        Scope.family(fid, actor, key),
        key,
        "export.create",
        body,
        build,
        lambda r: {"job_id": r["results"][0]["object_id"]},
    )
    return get_export(ctx, blobs, actor, fid, out["job_id"])


def get_export(
    ctx: AppContext, blobs: BlobStore, actor: Actor, fid: str, jid: str
) -> dict[str, Any]:
    require_member(load_membership(ctx, actor, fid))
    item = ctx.store.get(keys.family(fid), f"EXPORT#{jid}")
    if item is None:
        raise NotFound("导出任务不存在或已过期")
    out: dict[str, Any] = {
        "job_id": jid,
        "status": item["status"],
        "format": item["format"],
        "month": item["month"],
        "filters": item.get("filters", {}),
        "row_count": int(item["row_count"]),
    }
    if item["status"] == "ready":
        filename = quote(f"小家账本-{item['month']}.{'csv' if item['format'] == 'csv' else 'xlsx'}")
        out["download_url"] = blobs.presign_download(
            item["s3_key"], item["s3_version_id"], DOWNLOAD_TTL_S, filename
        )
        out["download_expires_at"] = keys.ts(ctx.clock() + timedelta(seconds=DOWNLOAD_TTL_S))
    return out
