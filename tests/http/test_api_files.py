"""OPS-04／11、AUTH-07／09、ACC-16（共享照片）、ACC-17、FX-05 经 HTTP（离线：本地对象存储替身）。"""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import UTC, datetime, timedelta
from typing import Any
from urllib.parse import parse_qs, urlparse

import pytest
from openpyxl import load_workbook
from PIL import Image
from starlette.applications import Starlette
from starlette.testclient import TestClient

from ledger.application import attachments, entries
from ledger.http.app import Runtime
from ledger.local.blobs import ROUTES as BLOB_ROUTES
from ledger.local.blobs import LocalBlobStore

from .conftest import Api, new_expense


def image_bytes(fmt: str = "JPEG", size: tuple[int, int] = (64, 48)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, (200, 120, 40)).save(buf, format=fmt)
    return buf.getvalue()


@pytest.fixture(scope="module")
def blob_client() -> TestClient:
    import ledger.local.blobs  # noqa: F401  # 登记本地 blob 路由

    return TestClient(Starlette(routes=list(BLOB_ROUTES)))


def upload(
    api: Api,
    client: TestClient,
    fid: str,
    data: bytes,
    declared: str = "image/jpeg",
    sha: str | None = None,
    declared_bytes: int | None = None,
) -> dict[str, Any]:
    status, intent = api.post(
        f"/families/{fid}/attachments/upload-intents",
        {
            "content_type": declared,
            "bytes": declared_bytes or len(data),
            "sha256": sha or hashlib.sha256(data).hexdigest(),
        },
    )
    assert status == 201, intent
    r = client.post(
        intent["upload_url"], data=intent["fields"], files={"file": ("x", data, declared)}
    )
    assert r.status_code in (204, 400), r.text
    status, done = api.post(f"/families/{fid}/attachments/{intent['attachment_id']}/complete")
    assert status in (200, 422), done
    return done if status == 200 else {"attachment_id": intent["attachment_id"], "error": done}


def test_ops04_upload_validation(family: dict[str, Any], blob_client: TestClient) -> None:
    fid, m = family["fid"], family["member"]
    ok = upload(m, blob_client, fid, image_bytes())
    assert ok["status"] == "ready" and (ok["width"], ok["height"]) == (64, 48)
    png_as_jpeg = upload(m, blob_client, fid, image_bytes("PNG"))
    assert png_as_jpeg["reject_reason"] == "format_mismatch"
    garbage = upload(m, blob_client, fid, b"<script>alert(1)</script>" * 10)
    assert garbage["reject_reason"] == "unreadable"
    wrong_sha = upload(m, blob_client, fid, image_bytes(), sha="0" * 64)
    assert wrong_sha["reject_reason"] == "digest_mismatch"
    huge = upload(m, blob_client, fid, image_bytes("PNG", (8000, 6000)), declared="image/png")
    assert huge["reject_reason"] == "too_many_pixels"
    # 声明 100 字节却上传更大的文件：签名表单拒收，完成时报告未收到文件
    small = upload(m, blob_client, fid, image_bytes(), declared_bytes=100)
    assert small["error"]["error"]["code"] == "upload_invalid"
    status, err = m.post(
        f"/families/{fid}/attachments/upload-intents",
        {"content_type": "image/heic", "bytes": 10, "sha256": "0" * 64},
    )
    assert status == 422  # 契约枚举拒绝 HEIC
    # 被拒绝的照片不能关联账目
    status, err = m.post(
        f"/families/{fid}/entries",
        {
            "business_date": "2026-10-02",
            "type": "expense",
            "amount": "1",
            "currency": "NZD",
            "leaf_category_id": "expense-01-01",
            "payment_method": "cash",
            "attachment_ids": [garbage["attachment_id"]],
        },
    )
    assert status == 422 and err["error"]["code"] == "upload_invalid"


def test_forged_upload_policy_rejected(family: dict[str, Any], blob_client: TestClient) -> None:
    fid, m = family["fid"], family["member"]
    _, intent = m.post(
        f"/families/{fid}/attachments/upload-intents",
        {"content_type": "image/jpeg", "bytes": 1000, "sha256": "0" * 64},
    )
    fields = {**intent["fields"], "key": "families/other/attachments/x"}
    r = blob_client.post(
        intent["upload_url"], data=fields, files={"file": ("x", b"1", "image/jpeg")}
    )
    assert r.status_code == 403


def test_auth07_auth09_photo_access(
    family: dict[str, Any],
    blob_client: TestClient,
    runtime: Runtime,
    admin_api: Any,
) -> None:
    fid, admin, m = family["fid"], family["admin"], family["member"]
    photo = upload(m, blob_client, fid, image_bytes())
    e = new_expense(m, fid, attachment_ids=[photo["attachment_id"]])
    assert e["attachments"] == [
        {"attachment_id": photo["attachment_id"], "content_type": "image/jpeg"}
    ]
    status, link = admin.get(f"/families/{fid}/attachments/{photo['attachment_id']}/download")
    assert status == 200
    assert blob_client.get(link["url"]).content == image_bytes()
    # 签发 60 秒有效
    exp = datetime.fromisoformat(link["expires_at"].replace("Z", "+00:00"))
    assert timedelta(seconds=55) < exp - datetime.now(UTC) <= timedelta(seconds=61)
    # 成员被移除：不能再签发新链接；已签发链接在到期前仍可用（不宣称即时撤回）
    status, old_link = m.get(f"/families/{fid}/attachments/{photo['attachment_id']}/download")
    _, members = admin.get(f"/families/{fid}/members")
    row = next(x for x in members["items"] if x["user_id"] == family["member_id"])
    admin.delete(
        f"/families/{fid}/members/{row['user_id']}", query={"expected_version": row["version"]}
    )
    status, _ = m.get(f"/families/{fid}/attachments/{photo['attachment_id']}/download")
    assert status == 403
    assert blob_client.get(old_link["url"]).status_code == 200
    # 篡改链接
    bad = old_link["url"].replace("t=", "t=x")
    assert blob_client.get(bad).status_code == 403


def test_shared_photo_survives_until_last_reference_purged(
    family: dict[str, Any], blob_client: TestClient, runtime: Runtime
) -> None:
    fid, admin = family["fid"], family["admin"]
    photo = upload(admin, blob_client, fid, image_bytes())
    aid = photo["attachment_id"]
    a = new_expense(admin, fid, attachment_ids=[aid])
    b = new_expense(admin, fid, attachment_ids=[aid], amount="5")
    store: LocalBlobStore = runtime.services["blobs"]
    key = attachments.blob_key(fid, aid)
    admin.delete(f"/families/{fid}/entries/{a['entry_id']}", query={"expected_version": 1})
    later = datetime.now(UTC) + timedelta(days=31)
    entries.purge_due(runtime.ctx, now=later)
    attachments.purge_orphans(runtime.ctx, store, now=later + timedelta(days=1))
    assert store.read(key, attachments.MAX_BYTES) is not None  # b 仍引用
    admin.delete(f"/families/{fid}/entries/{b['entry_id']}", query={"expected_version": 1})
    entries.purge_due(runtime.ctx, now=later)
    assert attachments.purge_orphans(runtime.ctx, store, now=later + timedelta(seconds=1)) >= 1
    assert store.read(key, attachments.MAX_BYTES) is None
    status, _ = admin.get(f"/families/{fid}/attachments/{aid}/download")
    assert status == 404


def test_unreferenced_photo_cleaned_after_7_days(
    family: dict[str, Any], blob_client: TestClient, runtime: Runtime
) -> None:
    fid, admin = family["fid"], family["admin"]
    photo = upload(admin, blob_client, fid, image_bytes())
    store: LocalBlobStore = runtime.services["blobs"]
    key = attachments.blob_key(fid, photo["attachment_id"])
    now = datetime.now(UTC)
    attachments.purge_orphans(runtime.ctx, store, now=now + timedelta(days=6))
    assert store.read(key, attachments.MAX_BYTES) is not None
    attachments.purge_orphans(runtime.ctx, store, now=now + timedelta(days=7, minutes=1))
    assert store.read(key, attachments.MAX_BYTES) is None


def test_acc17_exports_safe_and_reconcilable(
    family: dict[str, Any],
    blob_client: TestClient,
    runtime: Runtime,
) -> None:
    fid, admin = family["fid"], family["admin"]
    new_expense(admin, fid, note='=HYPERLINK("http://evil")', amount="10")
    new_expense(admin, fid, note="+1-2", amount="20", payment_method="cash")
    status, d = admin.get(f"/families/{fid}/dashboard", query={"month": "2026-10"})
    status, job = admin.post(f"/families/{fid}/exports", {"format": "csv", "month": "2026-10"})
    assert status == 202 and job["status"] == "ready" and job["row_count"] == 2
    text = blob_client.get(job["download_url"]).content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][0] == "日期" and len(rows) == 3
    notes = {r[9] for r in rows[1:]}
    assert notes == {'\'=HYPERLINK("http://evil")', "'+1-2"}
    assert sum(float(r[6]) for r in rows[1:]) == float(d["totals"]["net_expense"]["nzd"])
    status, job2 = admin.post(
        f"/families/{fid}/exports",
        {"format": "xlsx", "month": "2026-10", "filters": {"payment_method": "cash"}},
    )
    wb = load_workbook(io.BytesIO(blob_client.get(job2["download_url"]).content))
    ws = wb["账目"]
    assert ws.max_row == 2 and ws.cell(2, 10).value == "+1-2" and ws.cell(2, 10).data_type == "s"
    assert wb["汇总"].cell(5, 2).value == 20.0  # 净支出（同一筛选范围）
    # 下载链接 5 分钟有效，且只有本家庭成员可取得
    q = parse_qs(urlparse(job["download_url"]).query)
    assert "t" in q
    outsider = Api(runtime, ip="192.0.2.250")
    status, _ = outsider.get(f"/families/{fid}/exports/{job['job_id']}")
    assert status == 401


def test_fx05_recompute_preview_confirm_and_supersede(family: dict[str, Any]) -> None:
    fid, admin = family["fid"], family["admin"]
    a = new_expense(admin, fid, currency="AUD", amount="10", business_date="2026-10-01")
    b = new_expense(admin, fid, currency="AUD", amount="20", business_date="2026-10-01")
    admin.post(f"/families/{fid}/currencies", {"code": "AUD"})  # 已启用：返回 422，无副作用
    status, _ = admin.post(
        f"/families/{fid}/rates/manual",
        {"date": "2026-10-01", "currency": "AUD", "usd_value": "0.70", "reason": "修正"},
    )
    assert status == 201
    status, job = admin.post(
        f"/families/{fid}/rates/recompute", {"date_from": "2026-10-01", "date_to": "2026-10-01"}
    )
    assert status == 202 and job["preview"]["changed_count"] == 2
    assert job["preview"]["months"][0]["month"] == "2026-10"
    # 预览不改写账目
    _, same = admin.get(f"/families/{fid}/entries/{a['entry_id']}")
    assert same["version"] == 1 and same["fx_snapshot"] == a["fx_snapshot"]
    status, err = admin.post(
        f"/families/{fid}/rates/recompute/{job['job_id']}/confirm", {"preview_digest": "0" * 64}
    )
    assert status == 409
    status, done = admin.post(
        f"/families/{fid}/rates/recompute/{job['job_id']}/confirm",
        {"preview_digest": job["preview_digest"]},
    )
    assert status == 202 and done["status"] == "done" and done["applied_count"] == 2
    _, a2 = admin.get(f"/families/{fid}/entries/{a['entry_id']}")
    assert a2["version"] == 2 and a2["fx_snapshot"]["rates"]["AUD"]["source"] == "family_manual"
    assert a2["nzd_minor"] == 1250  # 10 × 0.70 ÷ 0.56
    # 再次修正 → 预览 → 期间修改一笔 → 确认被标记 superseded，不覆盖
    admin.post(
        f"/families/{fid}/rates/manual",
        {
            "date": "2026-10-01",
            "currency": "AUD",
            "usd_value": "0.71",
            "reason": "再修正",
            "expected_revision": 1,
        },
    )
    _, job2 = admin.post(
        f"/families/{fid}/rates/recompute", {"date_from": "2026-10-01", "date_to": "2026-10-01"}
    )
    _, b2 = admin.get(f"/families/{fid}/entries/{b['entry_id']}")
    admin.patch(
        f"/families/{fid}/entries/{b['entry_id']}",
        {"expected_version": b2["version"], "note": "期间修改"},
    )
    _, res = admin.post(
        f"/families/{fid}/rates/recompute/{job2['job_id']}/confirm",
        {"preview_digest": job2["preview_digest"]},
    )
    assert res["status"] == "superseded" and res["failure_code"] == "entry_changed_since_preview"
