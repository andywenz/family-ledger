"""历史 Excel 导入（离线，合成数据）：读取→预览→核对后导入→幂等→对账。"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from ledger.application import repo
from ledger.ops import history_import as hi

from .conftest import World

MAPPING = {
    "items": {
        "超市": {"type": "expense", "category": ["食品与餐饮", "超市综合购物"]},
        "餐厅": {"type": "expense", "category": ["食品与餐饮", "外出用餐"]},
        "亲戚给钱": {"type": "income", "category": ["他人赠予", "亲友赠予"]},
    },
    "channels": {"信用卡": "credit_card", "现金": "cash"},
}


def _workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "202609"
    ws.append(["日期", "事项", "币种", "金额", "合NZD", "合CNY", "渠道"])
    rows = [
        (dt.date(2026, 9, 21), "超市", "NZD", 40.3, 40.3, 161.2, "信用卡"),
        (dt.date(2026, 9, 22), "餐厅", "CNY", 88, 22, 88, "现金"),
        (dt.date(2026, 9, 23), "亲戚给钱", "NZD", -50, -50, -200, "现金"),  # 赠予→收入
        (dt.date(2026, 9, 24), "餐厅", "NZD", -12.5, -12.5, -50, "现金"),  # 负数支出：冲突
        (dt.date(2026, 9, 25), "神秘事项", "NZD", 9.9, 9.9, 39.6, "信用卡"),  # 未映射
        (dt.date(2026, 9, 26), "超市", "NZD", 15, 15, 60, "支付宝"),  # 未知方式
    ]
    for r in rows:
        ws.append(list(r))
    ws.append([None, "合计", None, None, "=SUM(E2:E7)", None, None])  # 非数据行
    ws2 = wb.create_sheet("202610")
    ws2.append(["日期", "事项", "币种", "金额", "合NZD", "合CNY", "渠道"])
    ws2.append([dt.date(2026, 10, 1), "超市", "NZD", 20, 20, 80, "信用卡"])
    out = path / "src.xlsx"
    wb.save(out)
    return out


def test_preview_flags_and_round_trip(tmp_path: Path) -> None:
    rows, skipped = hi.read_workbook(_workbook(tmp_path))
    assert len(rows) == 7 and skipped == [("202609", 8)]  # 汇总行跳过并列出
    prev = hi.build_preview(rows, MAPPING)
    by_item = {(p.note, p.business_date.day): p for p in prev}
    gift = by_item[("亲戚给钱", 23)]
    assert (gift.type, gift.amount, gift.include) == ("income", Decimal("50"), True)
    assert not by_item[("餐厅", 24)].include  # 负数不凭符号猜：默认不导入
    unknown = by_item[("神秘事项", 25)]
    assert (unknown.parent, unknown.leaf) == ("其他支出", "待分类") and unknown.hints
    assert by_item[("超市", 26)].method is None and by_item[("超市", 26)].hints
    path = tmp_path / "preview.xlsx"
    hi.write_preview(prev, path)
    back = hi.read_preview(path)
    assert [(p.sheet, p.row, p.amount, p.include, p.fingerprint) for p in back] == [
        (p.sheet, p.row, p.amount, p.include, p.fingerprint) for p in prev
    ]


def test_import_idempotent_and_reconciled(world: World, tmp_path: Path) -> None:
    w = world
    src = _workbook(tmp_path)
    rows, _ = hi.read_workbook(src)
    prev = hi.build_preview(rows, MAPPING)
    for p in prev:
        if p.method is None:
            p.method = "credit_card"  # 用户在预览表中补选方式
    batch = hi.file_digest(src)
    rep = hi.import_rows(w.ctx, w.admin, w.fid, prev, batch=batch)
    assert (rep.created, rep.skipped, rep.failed) == (6, 1, [])
    again = hi.import_rows(w.ctx, w.admin, w.fid, prev, batch=batch)
    assert (again.created, again.replayed) == (0, 6)  # 重跑命中原结果，不重复入账
    rec = hi.reconcile(w.ctx, w.fid, prev)
    assert rec["一致"], rec
    sep = [e for e in repo.month_entries(w.ctx, w.fid, "2026-09") if e.source == "import"]
    assert len(sep) == 5
    gift = next(e for e in sep if e.type == "income")
    assert gift.amount_minor == 5000 and "历史导入 202609!4" in gift.note
    cny = next(e for e in sep if e.currency == "CNY")
    assert "原表折算 NZD 22" in cny.note and cny.nzd_minor > 0  # 按消费日汇率重算，原折算留在备注


def test_missing_category_stops_before_writing(world: World, tmp_path: Path) -> None:
    w = world
    rows, _ = hi.read_workbook(_workbook(tmp_path))
    prev = hi.build_preview(rows, MAPPING)
    prev[0].leaf = "不存在的分类"
    with pytest.raises(ValueError, match="找不到这些分类"):
        hi.import_rows(w.ctx, w.admin, w.fid, prev, batch="b" * 64)
    assert [e for e in repo.month_entries(w.ctx, w.fid, "2026-09") if e.source == "import"] == []


def test_edited_preview_header_rejected(tmp_path: Path) -> None:
    rows, _ = hi.read_workbook(_workbook(tmp_path))
    path = tmp_path / "p.xlsx"
    hi.write_preview(hi.build_preview(rows, MAPPING), path)
    wb = openpyxl.load_workbook(path)
    wb.worksheets[0]["D1"] = "type"
    wb.save(path)
    with pytest.raises(ValueError, match="表头"):
        hi.read_preview(path)
